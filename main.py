import os
import psycopg2 as sq2
import json
import asyncio
import logging
import numpy as np
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters.command import Command
from aiogram.utils.keyboard import InlineKeyboardBuilder, ReplyKeyboardBuilder
from openai import OpenAI, OpenAIError
import tiktoken
import traceback
from dotenv import load_dotenv
from sentence_transformers import SentenceTransformer
from aiogram.exceptions import TelegramBadRequest
import datetime
import re
import time
from weasyprint import HTML  # Для генерации PDF
import tempfile
import PyPDF2
import docx
from pathlib import Path
import openpyxl  # Для Excel файлов
from pptx import Presentation  # Для PowerPoint файлов
import pandas as pd

# Загрузка переменных окружения из .env
load_dotenv()

# Получение конфигурации
SECRET_INVITE_TOKEN = os.getenv("SECRET_INVITE_TOKEN")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
BOT_TOKEN = os.getenv("BOT_TOKEN")
DB_HOST = os.getenv("DB_HOST")
DB_NAME = os.getenv("DB_NAME")
DB_USER = os.getenv("DB_USER")
DB_PASSWORD = os.getenv("DB_PASSWORD")
DB_PORT = os.getenv("DB_PORT")

# Инициализация моделей
client = OpenAI(api_key=OPENAI_API_KEY, base_url="https://api.deepseek.com")
bge_model = SentenceTransformer('BAAI/bge-m3')
max_token_count = 60000

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("bot_errors.log"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Словари для хранения состояния пользователей
user_system_prompts = {}  # user_id: prompt_text
chat_contexts = {}  # (chat_id, user_id): message_history
user_role_types = {}  # user_id: role_type (expert, product, marketing)
user_search_modes = {}  # user_id: search_mode (vector, context) - только для marketing
user_last_html = {}  # user_id: last_generated_html - для роли marketing
user_uploaded_files = {}  # user_id: {file_content: str, file_name: str, instructions: str}

def get_current_date():
    moscow_tz = datetime.timezone(datetime.timedelta(hours=3))
    return datetime.datetime.now(moscow_tz).strftime("%d-%m-%Y")

# Функция загрузки промта из БД по имени
def load_prompt_by_name(prompt_name):
    """Загружает системный промт из базы данных по имени"""
    try:
        conn = create_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT content FROM promt_data WHERE name = %s", (prompt_name,))
        result = cursor.fetchone()
        conn.close()
        if result:
            logger.info(f"Промт '{prompt_name}' успешно загружен из БД")
            return result[0]
        else:
            logger.error(f"Промт '{prompt_name}' не найден в базе данных")
            return None
    except Exception as e:
        logger.error(f"Ошибка загрузки промта '{prompt_name}': {str(e)}")
        return None

# Инициализация токенизатора
tokenizer = tiktoken.get_encoding("cl100k_base")

# Работа с базой
def create_db_connection():
    return sq2.connect(
        host=DB_HOST,
        dbname=DB_NAME,
        user=DB_USER,
        password=DB_PASSWORD,
        port=DB_PORT
    )

def init_access_table():
    """Создает таблицу для контроля доступа, если она не существует"""
    try:
        conn = create_db_connection()
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS access_control (
                user_id BIGINT PRIMARY KEY,
                username VARCHAR(255),
                full_name VARCHAR(255),
                is_allowed BOOLEAN DEFAULT FALSE
            )
        ''')
        conn.commit()
        conn.close()
    except Exception as e:
        logger.error(f"Ошибка инициализации таблица доступа: {str(e)}")

def get_user_access(user_id):
    """Проверяет наличие доступа у пользователя"""
    try:
        conn = create_db_connection()
        cursor = conn.cursor()
        cursor.execute('SELECT is_allowed FROM access_control WHERE user_id = %s', (user_id,))
        result = cursor.fetchone()
        conn.close()
        return result[0] if result else False
    except Exception as e:
        logger.error(f"Ошибка получения доступа пользователя {user_id}: {str(e)}")
        return False

def set_user_access(user_id, username, full_name, is_allowed=True):
    """Устанавливает доступ для пользователя"""
    try:
        conn = create_db_connection()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO access_control (user_id, username, full_name, is_allowed)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (user_id) DO UPDATE SET is_allowed = EXCLUDED.is_allowed
        ''', (user_id, username, full_name, is_allowed))
        conn.commit()
        conn.close()
    except Exception as e:
        logger.error(f"Ошибка установки доступа для пользователя {user_id}: {str(e)}")

async def vector_search(question_embedding, limit=100, role_type="expert"):
    """Выполняет векторный поиск в базе данных с учетом роли"""
    try:
        connection = create_db_connection()
        cursor = connection.cursor()
        base_sql = """
            SELECT TO_CHAR(date, 'YYYY-MM-DD') AS date_str, chunk_text
            FROM article_chunks
            {source_filter}
            ORDER BY embedding <-> %s::vector
            LIMIT %s
        """

        # Добавляем фильтр по source только для роли product
        if role_type == "product":
            source_filter = "WHERE source = 'product'"
            #limit = 12
        else:
            source_filter = ""
        
        sql = base_sql.format(source_filter=source_filter)

        cursor.execute(sql, (question_embedding.tolist(), limit))
        rows = cursor.fetchall()
        columns = [desc[0] for desc in cursor.description]
        result = []
        for row in rows:
            result.append(dict(zip(columns, row)))
        connection.close()
        return result
    except Exception as e:
        logger.error(f"Ошибка векторного поиска: {str(e)}")
        return []

# Подсчет токенов
def count_tokens(text):
    return len(tokenizer.encode(text))

# Обработка чанка данных с учетом контекста
async def process_data_chunk(chunk, chat_context, user_message, system_prompt):
    try:
        chunk_json = json.dumps(chunk, ensure_ascii=False)
        # Формируем сообщения с учетом контекста диалога
        messages = [{"role": "system", "content": system_prompt}]
        # Добавляем истории диалога
        for msg in chat_context:
            messages.append(msg)
        # Добавляем текущие данные и вопрос
        messages.append({
            "role": "user",
            "content": f"Контекст диалога: {chat_context}\n\nВопрос: {user_message}\nДанные: {chunk_json}"
        })
        response = client.chat.completions.create(
            model="deepseek-chat",
            max_tokens=1000,
            messages=messages,
            stream=False,
        )
        return response.choices[0].message.content
    except OpenAIError as ex:
        logger.error(f"Ошибка OpenAI при обработке чанка: {str(ex)}")
        return "Нет информации"
    except Exception as ex:
        logger.error(f"Ошибка обработки чанка: {str(ex)}")
        return "Нет информации"

# Форматирование контекста для ИИ
def format_chat_context(context, max_tokens=4000):
    try:
        formatted = []
        current_tokens = 0
        # Идем от последнего сообщения к первому
        for msg in reversed(context):
            msg_tokens = count_tokens(f"{msg['role']}: {msg['content']}")
            if current_tokens + msg_tokens > max_tokens:
                break
            formatted.insert(0, msg)  # Добавляем в начало для сохранения порядка
            current_tokens += msg_tokens
        return formatted
    except Exception as e:
        logger.error(f"Ошибка форматирования контекста: {str(e)}")
        return []

# Генерация эмбеддинга для текста
async def get_text_embedding(text):
    """Асинхронно генерирует векторное представление текста"""
    loop = asyncio.get_running_loop()
    # Выполняем в отдельном потоке, чтобы не блокировать event loop
    embedding = await loop.run_in_executor(None, bge_model.encode, text)
    return embedding / np.linalg.norm(embedding)  # Нормализация для L2-расстояния

# Обрезка ИИ для PDF
def trim_ai_answer(ai_answer):
    # Ищем начало <!DOCTYPE html>
    start_index = ai_answer.find('<!DOCTYPE html>')
    if start_index == -1:
        return ''
    
    # Обрезаем всё до <!DOCTYPE html>
    trimmed_start = ai_answer[start_index:]
    
    # Ищем последний </html> в обрезанной строке
    end_index = trimmed_start.rfind('</html>')
    if end_index != -1:
        # Обрезаем до конца тега </html> (длина тега — 7 символов)
        end_index += len('</html>')
        return trimmed_start[:end_index]
    
    return trimmed_start

# Генерация PDF из HTML
async def generate_pdf_from_html(html_content, user_id):
    """Генерирует PDF файл из HTML контента"""
    try:
        # Создаем уникальное имя файла
        timestamp = int(time.time())
        pdf_filename = f"report_{user_id}_{timestamp}.pdf"
        pdf_path = os.path.join("temp_pdfs", pdf_filename)
        # Создаем директорию, если ее нет
        os.makedirs("temp_pdfs", exist_ok=True)
        # Генерируем PDF
        HTML(string=html_content).write_pdf(pdf_path)
        logger.info(f"PDF сгенерирован: {pdf_path}")
        return pdf_path
    except Exception as e:
        logger.error(f"Ошибка генерации PDF: {str(e)}")
        raise

# Функция для чтения текста из файлов разных форматов
async def extract_text_from_file(file_path, file_type):
    """Извлекает текст из файла в зависимости от его типа"""
    try:
        if file_type == 'text/plain' or file_path.endswith('.txt'):
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
                
        elif file_type == 'application/pdf' or file_path.endswith('.pdf'):
            text = ""
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            return text
            
        elif (file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' or 
              file_path.endswith('.docx')):
            doc = docx.Document(file_path)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
        elif (file_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' or 
              file_path.endswith('.xlsx') or file_path.endswith('.xls')):
            # Обработка Excel файлов
            text = ""
            if file_path.endswith('.xlsx'):
                df = pd.read_excel(file_path, sheet_name=None)
            else:
                df = pd.read_excel(file_path, engine='xlrd', sheet_name=None)
            
            for sheet_name, data in df.items():
                text += f"--- Лист: {sheet_name} ---\n"
                text += data.to_string(index=False) + "\n\n"
            return text
            
        elif (file_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation' or 
              file_path.endswith('.pptx')):
            # Обработка PowerPoint файлов
            text = ""
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides):
                text += f"--- Слайд {i+1} ---\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
                text += "\n"
            return text
            
        else:
            logger.warning(f"Неизвестный тип файла: {file_type}")
            return None
            
    except Exception as e:
        logger.error(f"Ошибка чтения файла: {str(e)}")
        return None

# Основная функция обработки запроса
async def handle_user_query(user_message, message, chat_context, user_id, bot: Bot):
    try:
        # Получаем текущий промт пользователя
        current_prompt = user_system_prompts.get(user_id)
        if not current_prompt:
            # Если промт не установлен, используем эксперт по умолчанию
            current_prompt = load_prompt_by_name('expert_promt') or ""
            user_system_prompts[user_id] = current_prompt

        # Отправляем промежуточные сообщения с reply
        bot_instance = bot
        status_msg = await bot_instance.send_message(
            message.chat.id,
            '🔍 Анализирую вопрос...',
            reply_to_message_id=message.message_id
        )
        
        # Определяем режим поиска для роли marketing
        role_type = user_role_types.get(user_id, "expert")
        search_mode = user_search_modes.get(user_id, "vector")  # По умолчанию векторный поиск
        
        if role_type == "marketing" and search_mode == "context":
            # Режим работы только на контекста сообщения
            context_str = ""
            await bot_instance.edit_message_text(
                '💽 Формирую ответ на основе контекста...',
                chat_id=status_msg.chat.id,
                message_id=status_msg.message_id
            )
        else:
            # Векторный поиск для всех остальных случаев
            question_embedding = await get_text_embedding(user_message)
            await bot_instance.edit_message_text(
                '🔮 Ищу в базе знаний...',
                chat_id=status_msg.chat.id,
                message_id=status_msg.message_id
            )
            
            search_results = await vector_search(question_embedding, role_type=role_type)
            if not search_results:
                await bot_instance.edit_message_text(
                    "По вашему запросу ничего не найдено.",
                    chat_id=status_msg.chat.id,
                    message_id=status_msg.message_id
                )
                return "По вашему запросу ничего не найдено."

            # Форматирование результатов для ИИ
            formatted_results = []
            for i, result in enumerate(search_results, 1):
                formatted_results.append(
                    f"Результат {i}:\n"
                    f"Дата: {result['date_str']}\n"
                    f"Описание: {result['chunk_text']}\n"
                )
            context_str = "\n\n".join(formatted_results)

        # Формирование запроса к ИИ
        current_date = get_current_date()
        
        # Для роли marketing добавляем предыдущий HTML в контекст, если он есть
        messages = [
            {"role": "system", "content": f"Текущая системная дата: {current_date}\n{current_prompt}"},
            *format_chat_context(chat_context),
        ]
        
        # Добавляем информацию из загруженного файла, если он есть
        file_info = ""
        if role_type == "marketing" and user_id in user_uploaded_files:
            file_data = user_uploaded_files[user_id]
            file_info = f"""
                Загруженный файл: {file_data['file_name']}
                Инструкции по использованию файла: {file_data.get('instructions', 'Не указаны')}
                Содержимое файла: {file_data['file_content'][:2000]}... [файл обрезан для экономии токенов]
            """
        
        # Добавляем предыдущий HTML для роли marketing, если он есть
        if role_type == "marketing" and user_id in user_last_html and user_last_html[user_id]:
            messages.append({
                "role": "user", 
                "content": f"""
                    {file_info}
                    Предыдущая версия HTML: {user_last_html[user_id]}
                    Контекст диалога: {format_chat_context(chat_context)}
                    Релевантные данные: {context_str}
                    Вопрос: {user_message}
                    
                    Внеси правки в предыдущий HTML на основе нового запроса и информации из загруженного файла. Не генерируй полностью новый HTML, а модифицируй существующий.
                """
            })
        else:
            messages.append({
                "role": "user", 
                "content": f"""
                    {file_info}
                    Контекст диалога: {format_chat_context(chat_context)}
                    Релевантные данные: {context_str}
                    Вопрос: {user_message}
                """
            })

        # Запрос к ИИ для формирования ответа
        await bot_instance.edit_message_text(
            '💽 Формирую ответ...',
            chat_id=status_msg.chat.id,
            message_id=status_msg.message_id
        )

        # Определяем параметры генерации в зависимости от роли
        if role_type == "marketing":
            max_tokens = 8000
            temperature = 1.0
        elif role_type == "product":
            max_tokens = 1000
            temperature = 0.3
        elif role_type == "expert":
            max_tokens = 1000
            temperature = 0.3
        else:
            max_tokens = 1000
            temperature = 0.7    

        response = client.chat.completions.create(
            model = "deepseek-chat",
            max_tokens = max_tokens,
            temperature = temperature,
            messages = messages,
            stream = False,
        )
        result = response.choices[0].message.content

        # Для роли marketing сохраняем сгенерированный HTML
        if role_type == "marketing":
            user_last_html[user_id] = trim_ai_answer(result)

        # Удаляем статусное сообщение
        try:
            await bot_instance.delete_message(status_msg.chat.id, status_msg.message_id)
        except Exception:
            pass

        return result
    except Exception as e:
        logger.error(f"Ошибка обработки запроса: {str(e)}\n{traceback.format_exc()}")
        return "Произошла ошибка при обработке запроса. Пожалуйста, попробуйте еще раз."

# Функция очистки контекста пользователя
def clear_user_context(chat_id, user_id):
    """Очищает контекст диалога для указанного пользователя в чате"""
    context_key = (chat_id, user_id)
    if context_key in chat_contexts:
        del chat_contexts[context_key]
    # Также очищаем последний HTML для пользователя
    if user_id in user_last_html:
        del user_last_html[user_id]
    # Очищаем загруженные файлы
    if user_id in user_uploaded_files:
        del user_uploaded_files[user_id]
    logger.info(f"Контекст, HTML и файлы очищены для пользователя {user_id} в чате {chat_id}")
    return True

# Создание меню команд
def get_main_menu_keyboard():
    builder = ReplyKeyboardBuilder()
    builder.row(
        types.KeyboardButton(text="/role"),
        types.KeyboardButton(text="/clear"),
        types.KeyboardButton(text="/mode")
    )
    return builder.as_markup(resize_keyboard=True)

# Показ меню выбора роли
async def show_role_menu(chat_id, bot: Bot):
    builder = InlineKeyboardBuilder()
    builder.add(types.InlineKeyboardButton(
        text="Аналитик", callback_data="role_expert")
    )
    builder.add(types.InlineKeyboardButton(
        text="Продакт", callback_data="role_product")
    )
    builder.add(types.InlineKeyboardButton(
        text="Маркетинг", callback_data="role_marketing")
    )
    builder.adjust(2, 2)  # По две кнопки в строке
    await bot.send_message(
        chat_id,
        "👥 Выбери, кто тебе нужен?",
        reply_markup=builder.as_markup()
    )

# Показ меню выбора режима поиска для маркетинга
async def show_search_mode_menu(chat_id, bot: Bot):
    builder = InlineKeyboardBuilder()
    builder.add(types.InlineKeyboardButton(
        text="Поиск в БЗ", callback_data="mode_vector")
    )
    builder.add(types.InlineKeyboardButton(
        text="Только контекст", callback_data="mode_context")
    )
    builder.adjust(1)  # По одной кнопке в строке
    await bot.send_message(
        chat_id,
        "🔍 Выбери режим поиска для маркетинга:",
        reply_markup=builder.as_markup()
    )

# Обработчик выбора роли
async def select_role(callback: types.CallbackQuery, bot: Bot):
    user_id = callback.from_user.id
    chat_id = callback.message.chat.id
    role_type = callback.data.split("_")[1]
    
    # Соответствие callback-данных именам промтов в БД
    role_mapping = {
        "expert": "expert_promt",
        "product": "product_expert",
        "marketing": "pdf_gen"  # Новый тип роли
    }
    
    if role_type in role_mapping:
        prompt_name = role_mapping[role_type]
        new_prompt = load_prompt_by_name(prompt_name)
        if new_prompt:
            user_system_prompts[user_id] = new_prompt
            user_role_types[user_id] = role_type  # Сохраняем тип роли
            
            # Очищаем контекст перед сменой роли
            clear_user_context(chat_id, user_id)
            
            role_names = {
                "expert": "Аналитиком",
                "product": "Продактом",
                "marketing": "Маркетологом"  # Новое название роли
            }
            role_name = role_names[role_type]
            await callback.answer(f"Роль изменена на: {role_name}")
            await callback.message.edit_text(
                f"✅ Теперь ты общаешься с {role_name}!\n"
                f"Контекст предыдущего диалога очищен.\n\n"
                "Задавай новый вопрос или используй /role чтобы изменить роль"
            )
            
            # Для роли маркетинг показываем дополнительное меню выбора режима
            if role_type == "marketing":
                await show_search_mode_menu(chat_id, bot)
        else:
            await callback.answer("Ошибка загрузки роли. Попробуйте позже.")
    else:
        await callback.answer("Неизвестная роль")

# Обработчик выбора режима поиска
async def select_search_mode(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    mode_type = callback.data.split("_")[1]
    
    # Проверяем, что пользователь имеет роль маркетинг
    if user_role_types.get(user_id) != "marketing":
        await callback.answer("Режим поиска доступен только для роли Маркетинг")
        return
        
    # Сохраняем выбранный режим
    user_search_modes[user_id] = mode_type
    
    mode_names = {
        "vector": "Поиск в БЗ",
        "context": "Только контекст сообщения"
    }
    
    await callback.answer(f"Режим поиска изменен на: {mode_names[mode_type]}")
    await callback.message.edit_text(
        f"✅ Режим поиска изменен на: {mode_names[mode_type]}\n\n"
        "Теперь можно задавать вопросы или использовать /mode для изменения режима"
    )

# Хэндлер на команду /start
async def cmd_start(message: types.Message, bot: Bot):
    try:
        user_id = message.from_user.id
        args = message.text.split()
        bot_instance = bot
        
        # Проверка инвайт-токена
        if len(args) > 1 and args[1] == SECRET_INVITE_TOKEN:
            set_user_access(
                user_id=user_id,
                username=message.from_user.username,
                full_name=message.from_user.full_name,
                is_allowed=True
            )
            # Устанавливаем промт по умолчанию (эксперт)
            default_prompt = load_prompt_by_name('expert_promt') or ""
            user_system_prompts[user_id] = default_prompt
            user_role_types[user_id] = "expert"  # Роль по умолчанию
            
            await message.answer('''
👋 Привет! 🤖 Я твой личный AI-помощник
───────────────────────
Имею доступ к:
📈 Аналитике Альфа Инвестиций (прочитал более 1000 статей)
💼 Каталогу продуктов
⚡ Готовым инструментам

Просто запроси:
• "Справка по SBER"
• "Скрипт: [аналитика] + [продукт]"

👤 Используй роли (/role):
• Аналитик - даст экспертную справку по рынку или бумаге
• Продакт - даст исчерпывающую информацию по продукту
• Маркетинг - сделает тебе маркетинговый материал примерно за 3 мин. 

📁 Для роли Маркетинг ты можешь загружать файлы (TXT, PDF, DOCX, XLSX, PPTX) и указывать, как их использовать.

🆕 Хочешь начать с чистого листа?
- используй команду /clear                       

Только проверенные данные ✅
───────────────────────
✨ Экономь время. Укрепляй экспертизу. Бусти продажи!
''', reply_markup=get_main_menu_keyboard())
            
            # Предлагаем выбрать роль
            await show_role_menu(message.chat.id, bot_instance)
        else:
            # Проверяем есть ли доступ у пользователя
            if get_user_access(user_id):
                # Устанавливаем промт по умолчанию если еще не установлен
                if user_id not in user_system_prompts:
                    default_prompt = load_prompt_by_name('expert_promt') or ""
                    user_system_prompts[user_id] = default_prompt
                    user_role_types[user_id] = "expert"  # Роль по умолчанию
                
                await message.answer('''
🤖 С возвращением! Вы уже имеете доступ к боту. Можете задавать вопросы.
''', reply_markup=get_main_menu_keyboard())
                # Предлагаем выбрать роль
                await show_role_menu(message.chat.id, bot_instance)
            else:
                await message.answer("⛔ Доступ запрещен. Используйте правильную пригласительную ссылку.")
    except Exception as e:
        logger.error(f"Ошибка в обработке /start: {str(e)}\n{traceback.format_exc()}")

# Хэндлер на команду /role
async def cmd_role(message: types.Message, bot: Bot):
    user_id = message.from_user.id
    bot_instance = bot
    if get_user_access(user_id):
        await show_role_menu(message.chat.id, bot_instance)
    else:
        await message.answer("⛔ У вас нет доступа к этой команде")

# Хэндлер на команду /mode
async def cmd_mode(message: types.Message, bot: Bot):
    user_id = message.from_user.id
    if not get_user_access(user_id):
        await message.answer("⛔ У вас нет доступа к этой команде")
        return
        
    # Проверяем, что пользователь имеет роль маркетинг
    if user_role_types.get(user_id) != "marketing":
        await message.answer("Режим поиска доступен только для роли Маркетинг")
        return
        
    await show_search_mode_menu(message.chat.id, bot)

# Хэндлер на команду /clear
async def cmd_clear(message: types.Message):
    user_id = message.from_user.id
    chat_id = message.chat.id
    if not get_user_access(user_id):
        await message.answer("⛔ У вас нет доступа к этой команде")
        return
    
    # Очищаем контекст
    cleared = clear_user_context(chat_id, user_id)
    if cleared:
        await message.answer("✅ Контекст диалога очищен. Начните новый диалог.")
    else:
        await message.answer("ℹ️ Контекст уже был пуст. Можете начать новый диалог.")

# Обработчик загрузки файлов
async def handle_document(message: types.Message, bot: Bot):
    user_id = message.from_user.id
    
    # Проверяем, что пользователь имеет доступ и роль маркетинг
    if not get_user_access(user_id) or user_role_types.get(user_id) != "marketing":
        await message.answer("⛔ Загрузка файлов доступна только для роли Маркетинг с активным доступом")
        return
    
    try:
        # Получаем информацию о файле
        file_id = message.document.file_id
        file_name = message.document.file_name
        mime_type = message.document.mime_type
        
        # Проверяем поддерживаемые форматы
        supported_types = [
            'text/plain', 
            'application/pdf', 
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'application/vnd.ms-excel',
            'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        ]
        
        supported_extensions = ['.txt', '.pdf', '.docx', '.xlsx', '.xls', '.pptx']
        
        if (mime_type not in supported_types and 
            not any(file_name.lower().endswith(ext) for ext in supported_extensions)):
            await message.answer("❌ Поддерживаются только файлы TXT, PDF, DOCX, XLSX, XLS и PPTX форматов")
            return
        
        # Скачиваем файл
        file = await bot.get_file(file_id)
        file_path = file.file_path
        
        # Создаем временную директорию для загрузки
        os.makedirs("temp_uploads", exist_ok=True)
        download_path = f"temp_uploads/{user_id}_{file_name}"
        
        # Скачиваем файл
        await bot.download_file(file_path, download_path)
        
        # Извлекаем текст из файла
        file_content = await extract_text_from_file(download_path, mime_type)
        
        if not file_content:
            await message.answer("❌ Не удалось прочитать содержимое файла")
            return
        
        # Сохраняем информацию о файле
        user_uploaded_files[user_id] = {
            'file_name': file_name,
            'file_content': file_content,
            'instructions': message.caption or "Не указаны"
        }
        
        # Удаляем временный файл
        os.remove(download_path)
        
        await message.answer(
            f"✅ Файл '{file_name}' успешно загружен!\n"
            f"Инструкции: {message.caption or 'Не указаны'}\n\n"
            "Теперь задайте вопрос, и я учту информацию из файла при генерации ответа."
        )
        
    except Exception as e:
        logger.error(f"Ошибка обработки файла: {str(e)}\n{traceback.format_exc()}")
        await message.answer("❌ Произошла ошибка при обработке файла. Попробуйте еще раз.")

# Основной обработчик сообщений
async def send(message: types.Message, bot: Bot):
    try:
        # Пропускаем сообщения от ботов
        if message.from_user.is_bot:
            return
        
        # Пропускаем команды
        if message.text and message.text.startswith('/'):
            return
        
        user_id = message.from_user.id
        bot_instance = bot
        
        # Логируем полученное сообщение
        logger.info(f"Получено сообщение: chat_id={message.chat.id}, user_id={user_id}, text={message.text}")
        logger.info(f"Тип чата: {message.chat.type}")
        
        # Проверка доступа
        if not get_user_access(user_id):
            logger.info(f"Пользователь {user_id} не имеет доступа")
            # В группе не отвечаем, в личке предупреждаем
            if message.chat.type == "private":
                await message.answer("🔒 Доступ только по приглашению. Используйте пригласительную ссылку для активации.")
            return
        
        # Обработка групповых чатов
        text = message.text or ""
        is_private = message.chat.type == "private"
        bot_username = (await bot_instance.get_me()).username
        mentioned = False
        is_reply = False
        
        if not is_private:
            # Проверяем упоминание бота
            if bot_username:
                # Ищем упоминание в любом месте сообщения
                pattern = re.compile(fr'@?{re.escape(bot_username)}', re.IGNORECASE)
                if pattern.search(text):
                    mentioned = True
                    # Удаляем упоминание из текста
                    text = pattern.sub('', text, count=1).strip()
            
            # Проверяем, является ли сообщение ответом на сообщение бота
            if message.reply_to_message and message.reply_to_message.from_user.id == bot_instance.id:
                is_reply = True
            
            logger.info(f"Упоминание: {mentioned}, Ответ: {is_reply}")
            
            # Если нет ни упоминания, ни ответа, то игнорируем
            if not (mentioned or is_reply):
                logger.info("Сообщение не содержит упоминания или ответа на бота - игнорируем")
                return
        else:
            # В личных сообщениях обрабатываем все
            mentioned = True
        
        # Если после удаления упоминания текст пуст, игнорируем
        if not text:
            logger.info("Текст пуст после удаления упоминания - игнорируем")
            return
        
        logger.info(f"Обрабатываем сообщение: {text}")
        await asyncio.sleep(1)
        
        # Ключ контекста: (chat_id, user_id)
        context_key = (message.chat.id, user_id)
        
        # Инициализация контекста для нового пользователя/чата
        if context_key not in chat_contexts:
            chat_contexts[context_key] = []
        
        # Добавляем текущее сообщение пользователя в контекст
        user_message = {
            "role": "user",
            "content": text
        }
        chat_contexts[context_key].append(user_message)
        
        try:
            # Обработка запроса
            ai_answer = await handle_user_query(
                user_message=text,
                message=message,
                chat_context=chat_contexts[context_key],
                user_id=user_id,
                bot=bot
            )
            
            # Проверяем роль пользователя
            role_type = user_role_types.get(user_id, "expert")
            
            # Обработка для роли "Маркетинг"
            if role_type == "marketing":
                try:
                    # Генерируем PDF из HTML
                    fix_ai_answer = trim_ai_answer(ai_answer)
                    pdf_path = await generate_pdf_from_html(fix_ai_answer, user_id)
                    
                    # Отправляем PDF файл
                    with open(pdf_path, 'rb') as pdf_file:
                        await bot_instance.send_document(
                            chat_id=message.chat.id,
                            document=types.BufferedInputFile(
                                pdf_file.read(),
                                filename=os.path.basename(pdf_path)
                            )
                        )

                    # Удаляем временный файл
                    os.remove(pdf_path)
                    
                    # Добавляем служебное сообщение в контекст
                    ai_message = {
                        "role": "assistant",
                        "content": "[PDF документ сгенерирован и отправлен]"
                    }
                    chat_contexts[context_key].append(ai_message)
                    
                except Exception as e:
                    logger.error(f"Ошибка генерации PDF: {str(e)}")
                    error_msg = f"⚠️ Ошибка генерации PDF: {str(e)}"
                    await message.reply(error_msg)
                    
                    # Добавляем оригинальный ответ в контекст
                    ai_message = {
                        "role": "assistant",
                        "content": ai_answer
                    }
                    chat_contexts[context_key].append(ai_message)
                    
                    # Отправляем ответ как текст
                    await message.reply(f"Содержимое ответа:\n{ai_answer}")
            
            # Обработка для остальных ролей
            else:
                # Добавляем ответ ИИ в контекст
                ai_message = {
                    "role": "assistant",
                    "content": ai_answer
                }
                chat_contexts[context_key].append(ai_message)
                
                # Отправка ответа с reply
                try:
                    await message.reply(ai_answer, parse_mode="Markdown")
                except TelegramBadRequest:
                    await message.reply(ai_answer)  # Отправка без Markdown
        
        except Exception as ex:
            logger.error(f"Ошибка обработки запроса: {str(ex)}\n{traceback.format_exc()}")
            await message.reply("Ошибка в обработке данных. Попробуйте еще раз")
    except Exception as e:
        logger.error(f"Критическая ошибка в обработчике send: {str(e)}\n{traceback.format_exc()}")

# Запуск процесса поллинга
async def main():
    try:
        # Инициализация системы доступа
        init_access_table()
        
        # Создаем временные директории
        os.makedirs("temp_pdfs", exist_ok=True)
        os.makedirs("temp_uploads", exist_ok=True)
        
        # Инициализация бота и диспетчера
        bot = Bot(token=BOT_TOKEN)
        dp = Dispatcher()
        
        # Регистрация обработчиков
        dp.callback_query.register(select_role, F.data.startswith("role_"))
        dp.callback_query.register(select_search_mode, F.data.startswith("mode_"))
        dp.message.register(cmd_start, Command("start"), F.chat.type == "private")
        dp.message.register(cmd_role, Command("role"), F.chat.type == "private")
        dp.message.register(cmd_mode, Command("mode"), F.chat.type == "private")
        dp.message.register(cmd_clear, Command("clear"))
        dp.message.register(handle_document, F.document)
        dp.message.register(send)
        
        logger.info("Бот запущен")
        await dp.start_polling(bot)
    except Exception as e:
        logger.error(f"Критическая ошибка при запуске бота: {str(e)}\n{traceback.format_exc()}")

if __name__ == "__main__":
    asyncio.run(main())